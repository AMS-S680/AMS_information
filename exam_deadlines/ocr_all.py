import os
os.environ['OMP_NUM_THREADS']='1'
from pathlib import Path
import json, hashlib, re, concurrent.futures, time,sys
from pptx import Presentation
import pymupdf as fitz
import numpy as np
from PIL import Image,ImageOps
from pillow_heif import register_heif_opener
from rapidocr_onnxruntime import RapidOCR
ROOT=Path(__file__).resolve().parents[1]; OUT=ROOT/'exam_deadlines'
register_heif_opener()
engine=None
def init():
    global engine
    import rapidocr_onnxruntime.utils as ocr_utils
    import onnxruntime as ort
    import cv2
    cv2.setNumThreads(1)
    def options():
        opts=ort.SessionOptions(); opts.intra_op_num_threads=1; opts.inter_op_num_threads=1
        return opts
    ocr_utils.SessionOptions=options
    engine=RapidOCR()
def work(task):
    kind,source,key,page=task
    try:
        if kind=='pdf':
            with fitz.open(source) as d:
                p=d[page-1]; pix=p.get_pixmap(matrix=fitz.Matrix(min(1.6,1300/max(p.rect.width,p.rect.height)),min(1.6,1300/max(p.rect.width,p.rect.height))),alpha=False)
                im=Image.frombytes('RGB',(pix.width,pix.height),pix.samples)
        else:
            if str(source).lower().endswith('.wmf') and Path(str(source)+'.png').exists(): source=str(source)+'.png'
            im=ImageOps.exif_transpose(Image.open(source)).convert('RGB'); im.thumbnail((1800,1800))
        result,_=engine(np.array(im))
        text='\n'.join(x[1] for x in result) if result else ''
        return dict(key=key,page=page,text=text,confidence=float(np.mean([float(x[2]) for x in result])) if result else None,empty=not bool(text.strip()),error=None)
    except Exception as e: return dict(key=key,page=page,text='',empty=True,error=str(e))
def main():
    inventory=json.loads((OUT/'inventory.json').read_text(encoding='utf-8')); pages=json.loads((OUT/'pages.json').read_text(encoding='utf-8'))
    seen=set(); tasks=[]; aliases={}
    for item in inventory:
        key=item['key']; source=ROOT/item['source']
        if item['kind'] not in ['.pdf','.pptx','.jpg','.jpeg','.png','.heic']: continue
        if item['sha256'] in seen:
            aliases[key]=next(i['key'] for i in inventory if i['sha256']==item['sha256']); continue
        seen.add(item['sha256'])
        if item['kind']=='.pptx': source=OUT/'rendered_ppt'/f'{key}.pdf'
        if item['kind'] in ['.pdf','.pptx']:
            if not source.exists():
                prs=Presentation(ROOT/item['source'])
                for n,slide in enumerate(prs.slides,1):
                    def nested(shapes):
                        for sh in shapes:
                            if sh.shape_type==6:
                                yield from sh.shapes
                                yield from nested(sh.shapes)
                    top=list(slide.shapes)
                    extra=[]; shape_seen=set()
                    for sh in nested(slide.shapes):
                        if id(sh._element) not in shape_seen:
                            extra.append(sh); shape_seen.add(id(sh._element))
                    all_shapes=list(enumerate(top))+[(100+j,sh) for j,sh in enumerate(extra)]
                    hashes=set()
                    for j,shape in all_shapes:
                        if shape.shape_type==13:
                            try: image=shape.image
                            except Exception:
                                print(f'MISSING_LINKED_IMAGE {key} slide {n} shape {j}',flush=True)
                                continue
                            target=OUT/'images'/f'{key}_s{n}_{j}.{image.ext}'
                            target.write_bytes(image.blob)
                            hashes.add(hashlib.sha256(image.blob).hexdigest())
                            tasks.append(('image',str(target),key,n*1000+j))
                    for j,pic in enumerate(slide._element.xpath('.//p:pic')):
                        for blip in pic.xpath('.//a:blip'):
                            rid=blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                            if not rid: continue
                            part=slide.part.related_part(rid); blob=part.blob
                            digest=hashlib.sha256(blob).hexdigest()
                            if digest in hashes: continue
                            hashes.add(digest)
                            ext=Path(str(part.partname)).suffix
                            target=OUT/'images'/f'{key}_s{n}_xml{j}{ext}'
                            target.write_bytes(blob)
                            tasks.append(('image',str(target),key,n*1000+200+j))
                continue
            with fitz.open(source) as d:
                for n,p in enumerate(d,1):
                    if p.get_images() or len(p.get_text().strip())<80 or any(r['key']==key and r['page']==n and re.search(r'mid.?term|deadline|\bdue\b|timetable|assessment|schedule',r['text'],re.I) for r in pages):
                        tasks.append(('pdf',str(source),key,n))
        else: tasks.append(('image',str(source),key,1))
    (OUT/'ocr_aliases.json').write_text(json.dumps(aliases),encoding='utf-8')
    (OUT/'ocr_tasks.json').write_text(json.dumps(tasks,ensure_ascii=False),encoding='utf-8')
    done=set(); log=OUT/'ocr_pages.jsonl'
    if log.exists():
        for line in log.read_text(encoding='utf-8').splitlines():
            r=json.loads(line)
            if not r.get('error'): done.add((r['key'],r['page']))
    tasks=[t for t in tasks if (t[2],t[3]) not in done]
    print(f'OCR tasks: {len(tasks)}, already done: {len(done)}',flush=True)
    if '--prepare-only' in sys.argv: return
    with concurrent.futures.ProcessPoolExecutor(max_workers=8,initializer=init) as pool,log.open('a',encoding='utf-8') as f:
        for n,r in enumerate(pool.map(work,tasks,chunksize=1),1):
            f.write(json.dumps(r,ensure_ascii=False)+'\n'); f.flush()
            if n%25==0 or r.get('error'): print(f'{n}/{len(tasks)} '+json.dumps({k:v for k,v in r.items() if k not in ['text']},ensure_ascii=False),flush=True)
if __name__=='__main__': main()
